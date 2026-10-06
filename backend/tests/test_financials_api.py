import os
import subprocess
import sys
import uuid
from datetime import date, timedelta
from decimal import Decimal

import pytest
from sqlmodel import func, select

from app.core import finance
from app.core.config import settings
from app.models.financials import (
    ImmutableRecordError,
    PaymentPlan,
    PaymentReceipt,
    ReceiptCounter,
    StudentFee,
)
from tests.test_grades_api import GRC, Q, RC, publish_two_exams

API = "/api/v1/financials"
TODAY = date.today()
PAST = (TODAY - timedelta(days=40)).isoformat()
FUTURE = (TODAY + timedelta(days=60)).isoformat()


def D(value) -> Decimal:
    return Decimal(str(value))


def make_structure(client, world, who="admin", **over):
    body = dict(name="رسوم السنة", grade_level="الخامس", academic_year="2025-2026", total_amount="3000", max_discount_percent="20")
    body.update(over)
    r = client.post(f"{API}/fee-structures", json=body, headers=getattr(world, who))
    assert r.status_code == 201, r.text
    return r.json()


def make_plan(client, world, structure, who="admin", **over):
    body = dict(
        fee_structure_id=structure["id"], student_id=world.s1.id,
        schedule=dict(count=3, first_due_date=PAST, interval_months=3),  # الأول متأخر 40 يومًا، والباقي قادم
    )
    body.update(over)
    return client.post(f"{API}/plans", json=body, headers=getattr(world, who))


def statement(client, world, student_id, who="acc"):
    r = client.get(f"{API}/students/{student_id}/statement", headers=getattr(world, who))
    assert r.status_code == 200, r.text
    return r.json()


def pay(client, world, fee_id, amount, who="acc", **over):
    body = dict(student_fee_id=fee_id, amount=str(amount), method="cash")
    body.update(over)
    return client.post(f"{API}/receipts", json=body, headers=getattr(world, who))


@pytest.fixture()
def fee(client, world):
    """رسم 3000 بثلاثة أقساط 1000: الأول متأخر (قبل 40 يومًا) والبقية قادمة."""
    structure = make_structure(client, world)
    assert make_plan(client, world, structure).status_code == 201
    return statement(client, world, world.s1.id)["fees"][0]


# ------------------------------------------------------------- الرسوم والأقساط


def test_fee_structure_permissions_and_uniqueness(client, world):
    make_structure(client, world)
    assert client.post(f"{API}/fee-structures", json=dict(name="x1", grade_level="الخامس", academic_year="2025-2026", total_amount="1"), headers=world.acc).status_code == 403
    assert client.post(f"{API}/fee-structures", json=dict(name="رسوم السنة", grade_level="الخامس", academic_year="2025-2026", total_amount="5"), headers=world.admin).status_code == 409
    assert len(client.get(f"{API}/fee-structures", headers=world.acc).json()) == 1
    for who in ("ta", "p1"):
        assert client.get(f"{API}/fee-structures", headers=getattr(world, who)).status_code == 403
    for bad in ({"total_amount": "0"}, {"total_amount": "-5"}, {"max_discount_percent": "101"}, {"id": str(uuid.uuid4())}, {"created_by": 1}):
        body = {"name": "رسم", "grade_level": "الخامس", "academic_year": "2025-2026", "total_amount": "10", **bad}
        assert client.post(f"{API}/fee-structures", json=body, headers=world.admin).status_code == 422


def test_plan_generation_rounding_and_statuses(client, world):
    structure = make_structure(client, world, name="ألف", total_amount="1000")
    r = make_plan(client, world, structure, schedule=dict(count=3, first_due_date=FUTURE, interval_months=1))
    assert r.status_code == 201 and r.json()["created_student_ids"] == [world.s1.id]
    inst = statement(client, world, world.s1.id)["fees"][0]["installments"]
    assert [D(i["amount"]) for i in inst] == [D("333.33"), D("333.33"), D("333.34")]
    assert sum(D(i["amount"]) for i in inst) == D("1000")
    assert {i["status"] for i in inst} == {"upcoming"}


def test_plan_validation_and_permissions(client, world):
    structure = make_structure(client, world)
    explicit = lambda a, b: [dict(due_date=PAST, amount=a), dict(due_date=FUTURE, amount=b)]  # noqa: E731
    base = dict(fee_structure_id=structure["id"], student_id=world.s1.id)
    post = lambda who="admin", **kw: client.post(f"{API}/plans", json={**base, **kw}, headers=getattr(world, who))  # noqa: E731
    assert post(installments=explicit("1000", "1999")).status_code == 422  # لا يساوي 3000
    assert post(schedule=dict(count=2, first_due_date=PAST), installments=explicit("1500", "1500")).status_code == 422
    assert post().status_code == 422  # لا schedule ولا installments
    assert post(classroom_id=world.c1.id, schedule=dict(count=1, first_due_date=PAST)).status_code == 422  # طالب وفصل معًا
    assert post(schedule=dict(count=1, first_due_date=PAST), discount_percent="5").status_code == 422  # بلا سبب
    assert post(who="acc", schedule=dict(count=1, first_due_date=PAST), discount_percent="5", discount_reason="أخ").status_code == 403
    assert post(schedule=dict(count=1, first_due_date=PAST), discount_percent="21", discount_reason="أخ").status_code == 422  # > 20%
    for who in ("ta", "p1"):
        assert post(who=who, schedule=dict(count=1, first_due_date=PAST)).status_code == 403
    assert post(student_id=9999, schedule=dict(count=1, first_due_date=PAST)).status_code == 404
    # الدفع الصحيح بخصم من المدير
    r = post(installments=explicit("1200", "1200"), discount_percent="20", discount_reason="خصم أخ")
    assert r.status_code == 201, r.text
    fee = statement(client, world, world.s1.id)["fees"][0]
    assert (D(fee["amount_due"]), D(fee["discount_amount"]), D(fee["net_amount"]), D(fee["remaining_amount"])) == (D(3000), D(600), D(2400), D(2400))


def test_class_wide_plan_is_atomic_idempotent_and_grade_checked(client, world):
    structure = make_structure(client, world)
    body = dict(fee_structure_id=structure["id"], classroom_id=world.c1.id, schedule=dict(count=2, first_due_date=FUTURE))
    r = client.post(f"{API}/plans", json=body, headers=world.acc)
    assert r.status_code == 201 and sorted(r.json()["created_student_ids"]) == sorted([world.s1.id, world.s2.id])
    again = client.post(f"{API}/plans", json=body, headers=world.acc).json()
    assert again["created_student_ids"] == [] and len(again["skipped_student_ids"]) == 2
    other = make_structure(client, world, name="سادس", grade_level="السادس")
    assert client.post(f"{API}/plans", json={**body, "fee_structure_id": other["id"]}, headers=world.admin).status_code == 422


# ------------------------------------------------------------- الدفع

def test_payment_allocates_fifo_numbers_and_updates_balances(client, world, session):
    structure = make_structure(client, world)
    make_plan(client, world, structure)
    before = statement(client, world, world.s1.id)
    assert D(before["totals"]["overdue_amount"]) == 1000 and before["totals"]["overdue_count"] == 1
    assert [i["status"] for i in before["fees"][0]["installments"]] == ["overdue", "upcoming", "upcoming"]
    assert before["fees"][0]["installments"][0]["days_overdue"] >= 40
    fee = before["fees"][0]
    r1 = pay(client, world, fee["id"], "1500")
    assert r1.status_code == 201, r1.text
    receipt = r1.json()
    assert receipt["receipt_number"] == f"R-{TODAY.year}-000001" and receipt["kind"] == "payment"
    assert [D(a["amount"]) for a in receipt["allocations"]] == [D(1000), D(500)]
    assert receipt["collected_by"] == world.users["acc"].id and receipt["student_name"] == "طالب1"
    assert pay(client, world, fee["id"], "100").json()["receipt_number"].endswith("000002")
    st = statement(client, world, world.s1.id)
    f = st["fees"][0]
    assert (D(f["paid_amount"]), D(f["remaining_amount"])) == (D(1600), D(1400))
    assert [i["status"] for i in f["installments"]] == ["paid", "upcoming", "upcoming"]
    assert D(st["totals"]["overdue_amount"]) == 0 and st["totals"]["remaining_total"] == "1400.00"
    # السجل المخزَّن متسق مع الأقساط
    db_fee = session.get(StudentFee, uuid.UUID(fee["id"]))
    assert D(db_fee.paid_amount) == sum(D(p.paid_amount) for p in session.exec(select(PaymentPlan)).all())


def test_payment_rules(client, world):
    structure = make_structure(client, world)
    make_plan(client, world, structure)
    fee = statement(client, world, world.s1.id)["fees"][0]
    insts = fee["installments"]
    assert pay(client, world, fee["id"], "3000.01").status_code == 422  # أكثر من المتبقي
    assert pay(client, world, fee["id"], "0").status_code == 422
    assert pay(client, world, fee["id"], "-5").status_code == 422
    assert pay(client, world, fee["id"], "10", method="transfer").status_code == 422  # تحتاج مرجعًا
    assert pay(client, world, fee["id"], "10", method="transfer", reference="TRX-1").status_code == 201
    assert pay(client, world, fee["id"], "10", method="zain_cash", reference="ZC-9").status_code == 201
    assert pay(client, world, fee["id"], "10", method="bitcoin").status_code == 422
    assert pay(client, world, fee["id"], "10", paid_at=(TODAY + timedelta(days=1)).isoformat()).status_code == 422
    assert pay(client, world, fee["id"], "10", status="x").status_code == 422  # حقل زائد
    assert pay(client, world, str(uuid.uuid4()), "10").status_code == 404
    # قسط محدد: لا يتجاوز متبقيه، ويجب أن يكون من نفس الحساب
    assert pay(client, world, fee["id"], "1000", installment_id=insts[2]["id"]).status_code == 201
    assert pay(client, world, fee["id"], "1", installment_id=insts[2]["id"]).status_code == 422
    assert pay(client, world, fee["id"], "1", installment_id=str(uuid.uuid4())).status_code == 404
    for who in ("ta", "p1", "tb"):
        assert pay(client, world, fee["id"], "10", who=who).status_code == 403
    assert client.post(f"{API}/receipts", json={}, headers=world.acc).status_code == 422


def test_idempotency_key_prevents_double_charge(client, world):
    structure = make_structure(client, world)
    make_plan(client, world, structure)
    fee = statement(client, world, world.s1.id)["fees"][0]
    first = pay(client, world, fee["id"], "500", idempotency_key="key-12345678").json()
    second = pay(client, world, fee["id"], "500", idempotency_key="key-12345678").json()
    assert first["id"] == second["id"]
    assert D(statement(client, world, world.s1.id)["fees"][0]["paid_amount"]) == D(500)
    assert pay(client, world, fee["id"], "600", idempotency_key="key-12345678").status_code == 409


def test_payment_is_atomic_when_a_step_fails(session, world, monkeypatch, client):
    structure = make_structure(client, world)
    make_plan(client, world, structure)
    fee = statement(client, world, world.s1.id)["fees"][0]
    before = session.exec(select(func.count()).select_from(PaymentReceipt)).one()

    def boom(*args, **kwargs):
        raise RuntimeError("تعطّل بعد إنشاء السند وقبل تحديث الأقساط")

    monkeypatch.setattr(finance, "_apply_allocations", boom)
    with pytest.raises(RuntimeError):
        finance.record_payment(
            session, user=world.users["admin"], student_fee_id=uuid.UUID(fee["id"]), amount=D(500),
            method=finance.PaymentMethod.CASH, reference=None, paid_at=None, note=None,
            installment_id=None, idempotency_key=None,
        )
    session.expire_all()
    assert session.exec(select(func.count()).select_from(PaymentReceipt)).one() == before  # لا سند يتيم
    assert D(session.get(StudentFee, uuid.UUID(fee["id"])).paid_amount) == 0
    assert session.exec(select(func.count()).select_from(ReceiptCounter)).one() == 0  # ولا رقم محروق


# ------------------------------------------------------------- العكس وثبات السجلات


def test_reversal_flow(client, world, session):
    structure = make_structure(client, world)
    make_plan(client, world, structure)
    fee = statement(client, world, world.s1.id)["fees"][0]
    original = pay(client, world, fee["id"], "1500").json()
    path = f"{API}/receipts/{original['id']}/reverse"
    assert client.post(path, json={"reason": "خطأ في الإدخال"}, headers=world.acc).status_code == 403
    assert client.post(path, json={"reason": "x"}, headers=world.admin).status_code == 422  # سبب قصير
    r = client.post(path, json={"reason": "خطأ في الإدخال"}, headers=world.admin)
    assert r.status_code == 201, r.text
    reversal = r.json()
    assert reversal["kind"] == "reversal" and reversal["receipt_number"].startswith("V-")
    assert reversal["reversal_of_id"] == original["id"] and D(reversal["amount"]) == D(1500)
    f = statement(client, world, world.s1.id)["fees"][0]
    assert D(f["paid_amount"]) == 0 and all(D(i["paid_amount"]) == 0 for i in f["installments"])
    # السند الأصلي كما هو، ومعلَّم بأنه معكوس
    got = client.get(f"{API}/receipts/{original['id']}", headers=world.acc).json()
    assert got["is_reversed"] is True and D(got["amount"]) == D(1500) and got["reversal_reason"] is None
    # لا عكس مزدوج ولا عكس لسند عكس
    assert client.post(path, json={"reason": "مرة أخرى"}, headers=world.admin).status_code == 409
    assert client.post(f"{API}/receipts/{reversal['id']}/reverse", json={"reason": "عكس العكس"}, headers=world.admin).status_code == 422
    assert client.post(f"{API}/receipts/{uuid.uuid4()}/reverse", json={"reason": "غير موجود"}, headers=world.admin).status_code == 404
    # بعد العكس يمكن الدفع من جديد
    assert pay(client, world, fee["id"], "3000").status_code == 201


def test_receipts_are_immutable_in_orm(session, client, world):
    structure = make_structure(client, world)
    make_plan(client, world, structure)
    fee = statement(client, world, world.s1.id)["fees"][0]
    pay(client, world, fee["id"], "100")
    receipt = session.exec(select(PaymentReceipt)).one()
    receipt.amount = D(1)
    with pytest.raises(ImmutableRecordError):
        session.commit()
    session.rollback()
    session.delete(session.exec(select(PaymentReceipt)).one())
    with pytest.raises(ImmutableRecordError):
        session.commit()
    session.rollback()


def test_database_triggers_block_raw_updates_after_migration(tmp_path):
    """الـ migration يزرع Triggers تمنع UPDATE/DELETE حتى من SQL مباشر."""
    import sqlite3

    db = tmp_path / "m.db"
    env = {**os.environ, "DATABASE_URL": f"sqlite:///{db}"}
    root = os.path.dirname(os.path.dirname(__file__))
    subprocess.run([sys.executable, "-m", "alembic", "upgrade", "head"], cwd=root, env=env, check=True, capture_output=True)
    con = sqlite3.connect(db)
    names = {r[0] for r in con.execute("select name from sqlite_master where type='trigger'")}
    assert {"trg_payment_receipts_no_update", "trg_payment_receipts_no_delete", "trg_receipt_allocations_no_update", "trg_receipt_allocations_no_delete"} <= names
    con.execute("PRAGMA foreign_keys=OFF")
    con.execute(
        "insert into payment_receipts(id,receipt_number,kind,student_fee_id,student_id,amount,method,paid_at,created_at,updated_at,created_by) "
        "values ('a','R-1','PAYMENT','f',1,5,'CASH','2026-01-01','2026-01-01','2026-01-01',1)"
    )
    for statement_ in ("update payment_receipts set amount=1", "delete from payment_receipts"):
        with pytest.raises(sqlite3.DatabaseError, match="immutable"):
            con.execute(statement_)


# ------------------------------------------------------------- كشف الحساب والصلاحيات


def test_statement_visibility(client, world):
    structure = make_structure(client, world)
    make_plan(client, world, structure)
    fee = statement(client, world, world.s1.id)["fees"][0]
    pay(client, world, fee["id"], "700", method="transfer", reference="TRX-SECRET", note="ملاحظة داخلية")

    staff = statement(client, world, world.s1.id, who="acc")
    assert staff["receipts"][0]["collected_by_name"] and staff["receipts"][0]["reference"] == "TRX-SECRET"
    assert statement(client, world, world.s1.id, who="admin")["student_name"] == "طالب1"

    r = client.get(f"{API}/students/{world.s1.id}/statement", headers=world.p1)
    assert r.status_code == 200
    body = r.json()
    text = r.text
    assert D(body["totals"]["paid_total"]) == D(700) and body["receipts"][0]["amount"] == "700.00"
    for secret in ("TRX-SECRET", "ملاحظة داخلية", "collected_by", "created_by", "idempotency"):
        assert secret not in text
    # رمز التحقق + التفقيط لازمان لطباعة السند، ولا يكشفان شيئًا داخليًا
    assert body["receipts"][0]["amount_in_words"] == "سبعمائة دينار فقط لا غير" and body["receipts"][0]["verification_code"]
    # ولي أمر آخر / طالب غير مربوط / معلم / بلا دخول
    assert client.get(f"{API}/students/{world.s1.id}/statement", headers=world.p2).status_code == 404
    assert client.get(f"{API}/students/{world.s3.id}/statement", headers=world.p1).status_code == 404
    assert client.get(f"{API}/students/{world.s1.id}/statement", headers=world.ta).status_code == 403
    assert client.get(f"{API}/students/{world.s1.id}/statement").status_code == 401
    assert client.get(f"{API}/students/99999/statement", headers=world.acc).status_code == 404


def test_verify_receipt_code(client, world):
    structure = make_structure(client, world)
    make_plan(client, world, structure)
    fee = statement(client, world, world.s1.id)["fees"][0]
    receipt = pay(client, world, fee["id"], "250").json()
    ok = client.get(f"{API}/receipts/verify", params={"number": receipt["receipt_number"], "code": receipt["verification_code"]})
    assert ok.status_code == 200 and ok.json()["valid"] is True and D(ok.json()["amount"]) == D(250) and ok.json()["reversed"] is False
    bad = client.get(f"{API}/receipts/verify", params={"number": receipt["receipt_number"], "code": "0" * 20}).json()
    assert bad == {"valid": False, "receipt_number": None, "kind": None, "amount": None, "paid_at": None, "reversed": None}
    assert client.get(f"{API}/receipts/verify", params={"number": "R-2000-999999", "code": "x"}).json()["valid"] is False
    client.post(f"{API}/receipts/{receipt['id']}/reverse", json={"reason": "تصحيح"}, headers=world.admin)
    assert client.get(f"{API}/receipts/verify", params={"number": receipt["receipt_number"], "code": receipt["verification_code"]}).json()["reversed"] is True


# ------------------------------------------------------------- المتأخرون والملخص


def test_defaulters_report_and_summary(client, world):
    structure = make_structure(client, world)
    make_plan(client, world, structure)                       # s1: ثلاثة أقساط، أولها قبل 40 يومًا
    make_plan(client, world, structure, student_id=world.s2.id, schedule=dict(count=3, first_due_date=FUTURE))
    fee = statement(client, world, world.s1.id)["fees"][0]

    rows = client.get(f"{API}/defaulters", headers=world.acc).json()
    assert [r["student_id"] for r in rows] == [world.s1.id]                       # s2 كله قادم
    assert D(rows[0]["overdue_amount"]) == D(1000) and rows[0]["overdue_installments"] == 1
    assert rows[0]["days_overdue"] >= 40 and rows[0]["guardian_phone"] and rows[0]["financial_hold"] is True
    assert client.get(f"{API}/defaulters", params={"min_overdue": "1001"}, headers=world.admin).json() == []
    assert client.get(f"{API}/defaulters", params={"classroom_id": world.c2.id}, headers=world.admin).json() == []
    for who in ("ta", "p1"):
        assert client.get(f"{API}/defaulters", headers=getattr(world, who)).status_code == 403

    s = client.get(f"{API}/summary", headers=world.acc).json()
    assert (D(s["net_total"]), D(s["overdue_amount"]), s["overdue_students"], s["held_students"]) == (D(6000), D(1000), 1, 1)
    assert D(s["today_collected"]) == 0 and D(s["collection_rate"]) == 0

    pay(client, world, fee["id"], "1000")
    pay(client, world, fee["id"], "500", method="zain_cash", reference="ZC1")
    s = client.get(f"{API}/summary", headers=world.acc).json()
    assert D(s["today_collected"]) == D(1500) and s["today_receipts"] == 2 and D(s["month_collected"]) == D(1500)
    assert D(s["collection_rate"]) == D("25.00") and D(s["overdue_amount"]) == 0 and s["overdue_students"] == 0
    assert {m["method"]: D(m["total"]) for m in s["by_method_today"]} == {"cash": D(1000), "zain_cash": D(500)}
    # العكس يخصم من مقبوضات اليوم (صافي)
    first = client.get(f"{API}/receipts", headers=world.acc).json()
    cash = next(r for r in first if r["method"] == "cash")
    client.post(f"{API}/receipts/{cash['id']}/reverse", json={"reason": "خطأ تحصيل"}, headers=world.admin)
    assert D(client.get(f"{API}/summary", headers=world.acc).json()["today_collected"]) == D(500)
    assert len(client.get(f"{API}/receipts", params={"date_from": TODAY.isoformat(), "date_to": TODAY.isoformat()}, headers=world.acc).json()) == 3
    assert client.get(f"{API}/receipts", params={"date_to": (TODAY - timedelta(days=1)).isoformat()}, headers=world.acc).json() == []
    assert client.get(f"{API}/summary", headers=world.ta).status_code == 403


# ------------------------------------------------------------- الحجب المالي لشهادة ولي الأمر


def test_financial_hold_blocks_guardian_report_card(client, world, monkeypatch):
    publish_two_exams(client, world)
    for sid in (world.s1.id, world.s2.id):
        client.post(RC.format(sid) + "/publish", params=Q, headers=world.admin)
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 200  # لا ديون ⇒ لا حجب

    structure = make_structure(client, world)
    make_plan(client, world, structure)                                           # s1 متأخر بـ 1000
    make_plan(client, world, structure, student_id=world.s2.id, schedule=dict(count=3, first_due_date=FUTURE))
    fee = statement(client, world, world.s1.id)["fees"][0]

    blocked = client.get(GRC.format(world.s1.id), params=Q, headers=world.p1)
    assert blocked.status_code == 403
    assert blocked.json()["detail"] == {"code": "financial_hold", "message": "يرجى مراجعة الحسابات"}
    assert "68" not in blocked.text                                              # لا تسرّب لأي بيانات شهادة
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p2).status_code == 404   # ولي أمر آخر: لا يعرف شيئًا
    assert client.get(GRC.format(world.s2.id), params=Q, headers=world.p2).status_code == 200   # أقساطه قادمة فقط
    assert client.get(RC.format(world.s1.id), params=Q, headers=world.admin).status_code == 200  # الطاقم غير محجوب
    st = client.get(f"{API}/students/{world.s1.id}/statement", headers=world.p1).json()
    assert st["financial_hold"] is True and st["hold_message"] == "يرجى مراجعة الحسابات"

    # عتبة أعلى من المتأخر ⇒ لا حجب
    monkeypatch.setattr(settings, "FINANCIAL_HOLD_THRESHOLD_AMOUNT", D(1000))
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 200
    monkeypatch.setattr(settings, "FINANCIAL_HOLD_THRESHOLD_AMOUNT", D(999))
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 403
    # فترة سماح تغطي التأخر (40 يومًا)
    monkeypatch.setattr(settings, "FINANCIAL_HOLD_GRACE_DAYS", 45)
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 200
    monkeypatch.setattr(settings, "FINANCIAL_HOLD_GRACE_DAYS", 0)
    monkeypatch.setattr(settings, "FINANCIAL_HOLD_ENABLED", False)
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 200
    monkeypatch.setattr(settings, "FINANCIAL_HOLD_ENABLED", True)

    # السداد يرفع الحجب تلقائيًا، وعكس السند يعيده
    receipt = pay(client, world, fee["id"], "1000").json()
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 200
    client.post(f"{API}/receipts/{receipt['id']}/reverse", json={"reason": "سند خاطئ"}, headers=world.admin)
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 403


# ------------------------------------------------------------- تفقيط، استثناء الحجب، التصدير


@pytest.mark.parametrize(
    "amount, words",
    [
        ("150000", "مائة وخمسون ألف دينار فقط لا غير"),
        ("1", "دينار واحد فقط لا غير"),
        ("2", "ديناران فقط لا غير"),
        ("3", "ثلاثة دنانير فقط لا غير"),
        ("21", "واحد وعشرون دينارًا فقط لا غير"),
        ("100", "مائة دينار فقط لا غير"),
        ("1000", "ألف دينار فقط لا غير"),
        ("2000", "ألفا دينار فقط لا غير"),
        ("2500", "ألفان وخمسمائة دينار فقط لا غير"),
        ("3000", "ثلاثة آلاف دينار فقط لا غير"),
        ("11000", "أحد عشر ألف دينار فقط لا غير"),
        ("200000", "مائتا ألف دينار فقط لا غير"),
        ("2500500", "مليونان وخمسمائة ألف وخمسمائة دينار فقط لا غير"),
        ("1000000000", "مليار دينار فقط لا غير"),
        ("0.25", "مائتان وخمسون فلسًا فقط لا غير"),
        ("0", "صفر دينار فقط لا غير"),
    ],
)
def test_amount_in_words(amount, words):
    from app.core.arabic_numbers import amount_in_words

    assert amount_in_words(amount) == words
    with pytest.raises(ValueError):
        amount_in_words("-1")


def test_receipt_carries_amount_in_words(client, world):
    structure = make_structure(client, world, total_amount="300000")
    make_plan(client, world, structure)
    fee = statement(client, world, world.s1.id)["fees"][0]
    assert pay(client, world, fee["id"], "150000").json()["amount_in_words"] == "مائة وخمسون ألف دينار فقط لا غير"


def _overdue_setup(client, world):
    structure = make_structure(client, world)
    make_plan(client, world, structure)
    publish_two_exams(client, world)
    client.post(RC.format(world.s1.id) + "/publish", params=Q, headers=world.admin)
    return statement(client, world, world.s1.id)["fees"][0]


def test_hold_exemption_is_admin_only_and_lifts_the_hold(client, world):
    _overdue_setup(client, world)
    url = f"{API}/students/{world.s1.id}/hold-exemption"
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 403
    for who in ("acc", "ta", "p1"):
        assert client.put(url, json={"exempt": True, "notes": "اتفاق خاص مع الإدارة"}, headers=getattr(world, who)).status_code == 403
    assert client.put(url, json={"exempt": True}, headers=world.admin).status_code == 422          # السبب إلزامي
    assert client.put(url, json={"exempt": True, "notes": "قصير"}, headers=world.admin).status_code == 422
    assert client.put(url, json={"exempt": True, "notes": "x" * 501}, headers=world.admin).status_code == 422
    assert client.put(url, json={"exempt": True, "notes": "ok-ok", "id": 1}, headers=world.admin).status_code == 422
    assert client.put(f"{API}/students/9999/hold-exemption", json={"exempt": True, "notes": "اتفاق خاص"}, headers=world.admin).status_code == 404

    r = client.put(url, json={"exempt": True, "notes": "اتفاق خاص مع الإدارة"}, headers=world.admin)
    assert r.status_code == 200 and r.json()["hold_exempt"] is True
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 200      # الشهادة تُعرض رغم المتأخرات
    staff = statement(client, world, world.s1.id)
    assert staff["hold_exempt"] is True and staff["hold_exempt_notes"] == "اتفاق خاص مع الإدارة" and staff["financial_hold"] is False
    assert D(staff["totals"]["overdue_amount"]) == 1000                                              # الدين ما زال ظاهرًا
    row = client.get(f"{API}/defaulters", headers=world.acc).json()[0]
    assert row["hold_exempt"] is True and row["financial_hold"] is False
    assert client.get(f"{API}/summary", headers=world.acc).json()["held_students"] == 0
    guardian = client.get(f"{API}/students/{world.s1.id}/statement", headers=world.p1)
    assert "اتفاق خاص" not in guardian.text and "hold_exempt" not in guardian.text                   # لا تسرّب لولي الأمر
    cards = client.get(f"/api/v1/guardian/students/{world.s1.id}/report-cards", headers=world.p1).json()
    assert cards["financial_hold"] is False and len(cards["cards"]) == 1

    client.put(url, json={"exempt": False}, headers=world.admin)
    assert client.get(GRC.format(world.s1.id), params=Q, headers=world.p1).status_code == 403
    assert statement(client, world, world.s1.id)["hold_exempt_notes"] is None


def test_guardian_report_card_list_and_hold(client, world):
    _overdue_setup(client, world)
    url = f"/api/v1/guardian/students/{world.s1.id}/report-cards"
    held = client.get(url, headers=world.p1).json()
    assert held == {"financial_hold": True, "hold_message": "يرجى مراجعة الحسابات", "cards": []}
    assert client.get(url, headers=world.p2).status_code == 404       # ليس ابنه
    assert client.get(url, headers=world.acc).status_code == 403
    assert client.get(url).status_code == 401
    fee = statement(client, world, world.s1.id)["fees"][0]
    pay(client, world, fee["id"], "1000")
    ok = client.get(url, headers=world.p1).json()
    assert ok["financial_hold"] is False and ok["cards"][0]["term"] == "first" and "overall_percentage" not in ok["cards"][0]


def test_parent_sees_only_own_childrens_timetable_and_classes(client, world):
    # p1 ابنه في الفصل c1 فقط؛ جدول c2 وقائمة الفصول مقيَّدان
    assert client.get("/api/v1/timetable", params={"classroom_id": world.c1.id}, headers=world.p1).status_code == 200
    assert client.get("/api/v1/timetable", params={"classroom_id": world.c2.id}, headers=world.p1).status_code == 404
    names = [c["name"] for c in client.get("/api/v1/classes", headers=world.p1).json()]
    assert names == ["أ"]


def _load_xlsx(response):
    import io

    import openpyxl

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in response.headers["content-disposition"]
    return openpyxl.load_workbook(io.BytesIO(response.content)).active


def test_excel_exports(client, world, session):
    fee = _overdue_setup(client, world)
    make_plan(client, world, make_structure(client, world, name="ثانٍ"), student_id=world.s2.id)
    sheet = _load_xlsx(client.get(f"{API}/defaulters/export", headers=world.acc))
    assert sheet.sheet_view.rightToLeft is True
    assert [c.value for c in sheet[1]][:2] == ["الطالب", "الشعبة"]
    assert {sheet.cell(row=r, column=1).value for r in (2, 3)} == {"طالب1", "طالب2"}
    assert sheet.cell(row=4, column=1).value == "الإجمالي" and sheet.cell(row=4, column=5).value == 2000
    assert sheet.cell(row=2, column=5).number_format == "#,##0.##"
    assert _load_xlsx(client.get(f"{API}/defaulters/export", params={"classroom_id": world.c2.id}, headers=world.admin)).max_row == 2  # الترويسة + الإجمالي
    for who in ("ta", "p1"):
        assert client.get(f"{API}/defaulters/export", headers=getattr(world, who)).status_code == 403
        assert client.get(f"{API}/receipts/export", headers=getattr(world, who)).status_code == 403
    assert client.get(f"{API}/defaulters/export").status_code == 401

    # المقبوضات: عكس بقيمة سالبة والصافي في الأسفل، وحقن الصيغ يُخزَّن نصًا
    r1 = pay(client, world, fee["id"], "1000", note="=HYPERLINK(\"http://evil\",\"x\")").json()
    pay(client, world, fee["id"], "500", method="transfer", reference="+SUM(A1)")
    client.post(f"{API}/receipts/{r1['id']}/reverse", json={"reason": "-cmd|' /C calc'!A0"}, headers=world.admin)
    sheet = _load_xlsx(client.get(f"{API}/receipts/export", params={"date_from": TODAY.isoformat(), "date_to": TODAY.isoformat()}, headers=world.acc))
    values = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    assert [v[7] for v in values[:3]] == [1000, 500, -1000] and values[3][7] == 500 and values[3][0] == "الصافي"
    body = [(c.value, c.data_type) for row in sheet.iter_rows(min_row=2, max_row=4) for c in row if isinstance(c.value, str) and c.value[:1] in "=+-@"]
    assert body and all(dtype == "s" for _, dtype in body), body
    assert any("HYPERLINK" in v for v, _ in body) and any(v.startswith("+SUM") for v, _ in body)
    assert client.get(f"{API}/receipts/export", params={"date_to": (TODAY - timedelta(days=5)).isoformat()}, headers=world.acc).status_code == 200
