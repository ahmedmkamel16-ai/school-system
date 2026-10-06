import io
from datetime import date, datetime

import openpyxl
from fastapi import APIRouter, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlmodel import func, select

from app.core.audit import log_action
from app.core.deps import CurrentUser, SessionDep
from app.core.permissions import AcademicWritePermission
from app.core.scoping import visible_student_ids
from app.models.classroom import ClassRoom
from app.models.student import Student, StudentStatus
from app.schemas.student import (
    StudentCreate,
    StudentListResponse,
    StudentRead,
    StudentUpdate,
)

router = APIRouter(prefix="/students", tags=["students"])

EXPORT_HEADERS = [
    "الاسم الكامل",
    "الرقم الوطني",
    "تاريخ الميلاد",
    "الصف الدراسي",
    "الشعبة",
    "العنوان",
    "اسم ولي الأمر",
    "هاتف ولي الأمر",
]


def _to_read(student: Student, session: SessionDep) -> StudentRead:
    class_name = None
    if student.classroom_id is not None:
        classroom = session.get(ClassRoom, student.classroom_id)
        class_name = classroom.name if classroom else None
    return StudentRead(**student.model_dump(), class_name=class_name)


def _ensure_visible(student_id: int, current_user, session: SessionDep) -> None:
    visible_ids = visible_student_ids(current_user, session)
    if visible_ids is not None and student_id not in visible_ids:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="الطالب غير موجود")


def _build_query(
    current_user,
    session: SessionDep,
    grade_level: str | None,
    classroom_id: int | None,
    student_status: StudentStatus | None,
    search: str | None,
):
    visible_ids = visible_student_ids(current_user, session)
    query = select(Student)
    if visible_ids is not None:
        query = query.where(Student.id.in_(visible_ids))
    if grade_level:
        query = query.where(Student.grade_level == grade_level)
    if classroom_id:
        query = query.where(Student.classroom_id == classroom_id)
    if student_status:
        query = query.where(Student.status == student_status)
    if search:
        pattern = f"%{search}%"
        query = query.where(
            (Student.full_name.like(pattern)) | (Student.national_id.like(pattern))
        )
    return query, visible_ids


@router.get("", response_model=StudentListResponse)
def list_students(
    session: SessionDep,
    current_user: CurrentUser,
    grade_level: str | None = Query(default=None),
    classroom_id: int | None = Query(default=None),
    student_status: StudentStatus | None = Query(default=None, alias="status"),
    search: str | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=25, ge=1, le=200),
) -> StudentListResponse:
    query, visible_ids = _build_query(
        current_user, session, grade_level, classroom_id, student_status, search
    )
    if visible_ids is not None and not visible_ids:
        return StudentListResponse(items=[], total=0)

    total = session.exec(
        select(func.count()).select_from(query.with_only_columns(Student.id).subquery())
    ).one()

    students = session.exec(
        query.order_by(Student.id).offset((page - 1) * page_size).limit(page_size)
    ).all()
    return StudentListResponse(
        items=[_to_read(student, session) for student in students], total=total
    )


@router.get("/export")
def export_students(
    session: SessionDep,
    current_user: CurrentUser,
    grade_level: str | None = Query(default=None),
    classroom_id: int | None = Query(default=None),
    student_status: StudentStatus | None = Query(default=None, alias="status"),
    search: str | None = Query(default=None),
) -> StreamingResponse:
    query, visible_ids = _build_query(
        current_user, session, grade_level, classroom_id, student_status, search
    )
    students = [] if (visible_ids is not None and not visible_ids) else session.exec(
        query.order_by(Student.id)
    ).all()

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "الطلاب"
    sheet.append(EXPORT_HEADERS)
    for student in students:
        classroom = (
            session.get(ClassRoom, student.classroom_id)
            if student.classroom_id
            else None
        )
        sheet.append(
            [
                student.full_name,
                student.national_id,
                student.birth_date.isoformat(),
                student.grade_level,
                classroom.name if classroom else "",
                student.address or "",
                student.guardian_name,
                student.guardian_phone,
            ]
        )
    sheet.sheet_view.rightToLeft = True

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=students.xlsx"},
    )


@router.post("/import")
def import_students(
    session: SessionDep, current_user: AcademicWritePermission, file: UploadFile
) -> dict:
    workbook = openpyxl.load_workbook(io.BytesIO(file.file.read()))
    sheet = workbook.active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    created = 0
    errors: list[str] = []

    for index, row in enumerate(rows, start=2):
        if not row or not row[0]:
            continue
        try:
            full_name, national_id, birth_date_raw, grade_level, classroom_name, address, guardian_name, guardian_phone = (
                (list(row) + [None] * 8)[:8]
            )
            if not all([full_name, national_id, birth_date_raw, grade_level, guardian_name, guardian_phone]):
                errors.append(f"صف {index}: بيانات ناقصة")
                continue
            if session.exec(
                select(Student).where(Student.national_id == str(national_id))
            ).first():
                errors.append(f"صف {index}: الرقم الوطني {national_id} مسجل بالفعل")
                continue

            if isinstance(birth_date_raw, datetime):
                birth_date_value = birth_date_raw.date()
            elif isinstance(birth_date_raw, date):
                birth_date_value = birth_date_raw
            else:
                birth_date_value = date.fromisoformat(str(birth_date_raw))

            classroom_id = None
            if classroom_name:
                classroom = session.exec(
                    select(ClassRoom).where(
                        ClassRoom.name == str(classroom_name),
                        ClassRoom.grade_level == str(grade_level),
                    )
                ).first()
                classroom_id = classroom.id if classroom else None

            student = Student(
                full_name=str(full_name),
                national_id=str(national_id),
                birth_date=birth_date_value,
                grade_level=str(grade_level),
                address=str(address) if address else None,
                guardian_name=str(guardian_name),
                guardian_phone=str(guardian_phone),
                classroom_id=classroom_id,
            )
            session.add(student)
            session.commit()
            created += 1
        except Exception as exc:  # noqa: BLE001
            session.rollback()
            errors.append(f"صف {index}: خطأ غير متوقع ({exc})")

    if created:
        log_action(
            session,
            current_user,
            "import",
            "student",
            None,
            f"استورد {created} طالبًا من ملف Excel",
        )

    return {"created": created, "errors": errors}


@router.post("", response_model=StudentRead, status_code=status.HTTP_201_CREATED)
def create_student(
    student_in: StudentCreate, session: SessionDep, current_user: AcademicWritePermission
) -> StudentRead:
    existing = session.exec(
        select(Student).where(Student.national_id == student_in.national_id)
    ).first()
    if existing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="الرقم الوطني مسجل بالفعل",
        )
    student = Student.model_validate(student_in)
    session.add(student)
    session.commit()
    session.refresh(student)
    log_action(
        session, current_user, "create", "student", student.id,
        f"أضاف الطالب {student.full_name}",
    )
    return _to_read(student, session)


@router.post("/bulk-move")
def bulk_move_students(
    payload: dict, session: SessionDep, current_user: AcademicWritePermission
) -> dict:
    student_ids: list[int] = payload.get("student_ids", [])
    classroom_id: int | None = payload.get("classroom_id")
    grade_level: str | None = payload.get("grade_level")
    if not student_ids:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="لم يتم تحديد أي طالب")

    updated = 0
    for student_id in student_ids:
        student = session.get(Student, student_id)
        if not student:
            continue
        if classroom_id is not None:
            student.classroom_id = classroom_id
            if not grade_level:
                classroom = session.get(ClassRoom, classroom_id)
                if classroom:
                    student.grade_level = classroom.grade_level
        if grade_level:
            student.grade_level = grade_level
        session.add(student)
        updated += 1
    session.commit()
    log_action(
        session, current_user, "bulk_move", "student", None,
        f"نقل {updated} طالبًا دفعة واحدة",
    )
    return {"updated": updated}


@router.get("/{student_id}", response_model=StudentRead)
def get_student(student_id: int, session: SessionDep, current_user: CurrentUser) -> StudentRead:
    _ensure_visible(student_id, current_user, session)
    student = session.get(Student, student_id)
    if not student:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="الطالب غير موجود")
    return _to_read(student, session)


@router.patch("/{student_id}", response_model=StudentRead)
def update_student(
    student_id: int, student_in: StudentUpdate, session: SessionDep, current_user: AcademicWritePermission
) -> StudentRead:
    student = session.get(Student, student_id)
    if not student:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="الطالب غير موجود")
    for field, value in student_in.model_dump(exclude_unset=True).items():
        setattr(student, field, value)
    session.add(student)
    session.commit()
    session.refresh(student)
    log_action(
        session, current_user, "update", "student", student.id,
        f"عدّل بيانات الطالب {student.full_name}",
    )
    return _to_read(student, session)


@router.delete("/{student_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_student(student_id: int, session: SessionDep, current_user: AcademicWritePermission) -> None:
    student = session.get(Student, student_id)
    if not student:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="الطالب غير موجود")
    full_name = student.full_name
    session.delete(student)
    session.commit()
    log_action(session, current_user, "delete", "student", student_id, f"حذف الطالب {full_name}")
